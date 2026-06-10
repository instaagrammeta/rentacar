"""Reporting service (Module 11) with professional Excel & PDF exports.

Excel exports are produced with OpenPyXL and feature styled green headers, the
company logo, auto-filters and auto-sized columns. The same datasets can be
exported to PDF through :mod:`app.services.pdf_service` helpers.
"""
from __future__ import annotations

from datetime import date, datetime, timedelta
from pathlib import Path

from openpyxl import Workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet
from sqlalchemy import func, select

from flask import current_app

from app.extensions import db
from app.models import Car, Client, Payment, Rental
from app.models.enums import RentalStatus
from app.repositories import settings_repo

HEADER_FILL = PatternFill(start_color="1B5E20", end_color="1B5E20", fill_type="solid")
HEADER_FONT = Font(color="FFFFFF", bold=True, size=11)
TITLE_FONT = Font(color="1B5E20", bold=True, size=16)
THIN_SIDE = Side(style="thin", color="C8E6C9")
BORDER = Border(left=THIN_SIDE, right=THIN_SIDE, top=THIN_SIDE, bottom=THIN_SIDE)


def _export_path(name: str) -> Path:
    export_dir: Path = current_app.config["EXPORT_DIR"] / "reports"
    export_dir.mkdir(parents=True, exist_ok=True)
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    return export_dir / f"{name}_{timestamp}.xlsx"


def _style_sheet(ws: Worksheet, title: str, headers: list[str], rows: list[list]) -> None:
    """Apply the standard styled layout to a worksheet."""
    settings = settings_repo.get_settings()

    # Optional logo in the top-left corner.
    start_row = 1
    if settings.logo_path:
        logo = current_app.config["UPLOAD_DIR"] / settings.logo_path
        if logo.exists():
            try:
                img = XLImage(str(logo))
                img.width, img.height = 120, 60
                ws.add_image(img, "A1")
                start_row = 4
            except Exception:  # pragma: no cover
                start_row = 1

    # Company name + report title.
    ws.cell(row=start_row, column=1, value=settings.company_name).font = TITLE_FONT
    ws.cell(row=start_row + 1, column=1, value=title).font = Font(bold=True, size=12)
    ws.cell(
        row=start_row + 2,
        column=1,
        value=f"Сформировано: {datetime.now().strftime('%d.%m.%Y %H:%M')}",
    ).font = Font(italic=True, size=9, color="666666")

    header_row = start_row + 4
    for col, header in enumerate(headers, start=1):
        cell = ws.cell(row=header_row, column=col, value=header)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = BORDER

    for r, row in enumerate(rows, start=header_row + 1):
        for c, value in enumerate(row, start=1):
            cell = ws.cell(row=r, column=c, value=value)
            cell.border = BORDER
            cell.alignment = Alignment(vertical="center")

    # Auto filter over the header + data range.
    last_col = get_column_letter(len(headers))
    last_row = header_row + len(rows)
    ws.auto_filter.ref = f"A{header_row}:{last_col}{max(header_row, last_row)}"

    # Auto-size columns based on the longest cell value.
    for col in range(1, len(headers) + 1):
        letter = get_column_letter(col)
        max_len = len(str(headers[col - 1]))
        for row in rows:
            value = row[col - 1] if col - 1 < len(row) else ""
            max_len = max(max_len, len(str(value)) if value is not None else 0)
        ws.column_dimensions[letter].width = min(max_len + 4, 50)

    ws.freeze_panes = ws.cell(row=header_row + 1, column=1)


def _save(wb: Workbook, name: str) -> str:
    path = _export_path(name)
    wb.save(path)
    return str(path)


# --------------------------------------------------------------------- queries
def _revenue_between(start: datetime, end: datetime) -> float:
    stmt = select(func.coalesce(func.sum(Payment.amount), 0.0)).where(
        Payment.paid_at >= start, Payment.paid_at < end
    )
    return float(db.session.execute(stmt).scalar_one())


def daily_revenue(day: date | None = None) -> dict:
    day = day or date.today()
    start = datetime.combine(day, datetime.min.time())
    payments = db.session.execute(
        select(Payment).where(Payment.paid_at >= start, Payment.paid_at < start + timedelta(days=1))
    ).scalars().all()
    total = sum(p.amount for p in payments)
    return {
        "date": day.isoformat(),
        "total": round(total, 2),
        "payments": [p.to_dict() for p in payments],
    }


def monthly_revenue(year: int, month: int) -> dict:
    start = datetime(year, month, 1)
    end = (start.replace(day=28) + timedelta(days=4)).replace(day=1)
    total = _revenue_between(start, end)
    return {"year": year, "month": month, "total": round(total, 2)}


def yearly_revenue(year: int) -> dict:
    months = []
    for m in range(1, 13):
        months.append(monthly_revenue(year, m))
    total = sum(m["total"] for m in months)
    return {"year": year, "total": round(total, 2), "months": months}


def most_profitable_cars(limit: int = 10) -> list[dict]:
    stmt = (
        select(Car, func.coalesce(func.sum(Payment.amount), 0.0).label("revenue"))
        .join(Rental, Rental.car_id == Car.id)
        .join(Payment, Payment.rental_id == Rental.id)
        .group_by(Car.id)
        .order_by(func.coalesce(func.sum(Payment.amount), 0.0).desc())
        .limit(limit)
    )
    rows = db.session.execute(stmt).all()
    return [{"car": car.display_name, "revenue": round(float(rev), 2)} for car, rev in rows]


def active_rentals() -> list[dict]:
    rentals = db.session.execute(
        select(Rental).where(Rental.status == RentalStatus.ACTIVE)
    ).scalars().all()
    return [r.to_dict() for r in rentals]


def debtors() -> list[dict]:
    """Clients whose completed returns left a positive final payment owed."""
    rentals = db.session.execute(
        select(Rental).where(Rental.status == RentalStatus.COMPLETED)
    ).scalars().all()
    result = []
    for rental in rentals:
        ret = rental.vehicle_return
        if ret and ret.final_payment > 0:
            result.append(
                {
                    "client": rental.client.full_name if rental.client else "—",
                    "contract_number": rental.contract_number,
                    "amount": round(ret.final_payment, 2),
                }
            )
    return result


def client_statistics() -> list[dict]:
    stmt = (
        select(Client, func.count(Rental.id).label("rentals"))
        .outerjoin(Rental, Rental.client_id == Client.id)
        .group_by(Client.id)
        .order_by(func.count(Rental.id).desc())
    )
    rows = db.session.execute(stmt).all()
    return [
        {"client": client.full_name, "code": client.client_code, "rentals": count}
        for client, count in rows
    ]


# ------------------------------------------------------------------- exporters
def export_report_excel(report_type: str, params: dict | None = None) -> str:
    """Generate a styled Excel report and return the saved file path."""
    params = params or {}
    wb = Workbook()
    ws = wb.active

    if report_type == "daily_revenue":
        data = daily_revenue(
            date.fromisoformat(params["date"]) if params.get("date") else None
        )
        ws.title = "Выручка за день"
        headers = ["Квитанция", "Клиент", "Тип", "Способ", "Сумма", "Дата"]
        rows = [
            [
                p["receipt_number"],
                p.get("client_name", ""),
                p.get("payment_type_label", ""),
                p.get("method_label", ""),
                p["amount"],
                p["paid_at"],
            ]
            for p in data["payments"]
        ]
        rows.append(["", "", "", "ИТОГО", data["total"], ""])
        _style_sheet(ws, f"Выручка за {data['date']}", headers, rows)

    elif report_type == "most_profitable_cars":
        data = most_profitable_cars(params.get("limit", 10))
        ws.title = "Прибыльные авто"
        _style_sheet(
            ws,
            "Самые прибыльные автомобили",
            ["Автомобиль", "Выручка"],
            [[d["car"], d["revenue"]] for d in data],
        )

    elif report_type == "active_rentals":
        data = active_rentals()
        ws.title = "Активные аренды"
        _style_sheet(
            ws,
            "Активные аренды",
            ["Договор", "Клиент", "Автомобиль", "Начало", "Окончание", "Сумма"],
            [
                [
                    r["contract_number"],
                    r.get("client_name", ""),
                    r.get("car_name", ""),
                    r["rental_start"],
                    r["rental_end"],
                    r["total_price"],
                ]
                for r in data
            ],
        )

    elif report_type == "debtors":
        data = debtors()
        ws.title = "Должники"
        _style_sheet(
            ws,
            "Список должников",
            ["Клиент", "Договор", "Сумма долга"],
            [[d["client"], d["contract_number"], d["amount"]] for d in data],
        )

    elif report_type == "client_statistics":
        data = client_statistics()
        ws.title = "Статистика клиентов"
        _style_sheet(
            ws,
            "Статистика по клиентам",
            ["Клиент", "Код", "Кол-во аренд"],
            [[d["client"], d["code"], d["rentals"]] for d in data],
        )

    elif report_type == "yearly_revenue":
        year = int(params.get("year", date.today().year))
        data = yearly_revenue(year)
        ws.title = "Годовая выручка"
        rows = [[m["month"], m["total"]] for m in data["months"]]
        rows.append(["ИТОГО", data["total"]])
        _style_sheet(ws, f"Выручка за {year} год", ["Месяц", "Выручка"], rows)

    else:
        ws.title = "Отчёт"
        _style_sheet(ws, "Неизвестный отчёт", ["Сообщение"], [["Тип отчёта не поддерживается"]])

    return _save(wb, report_type)
